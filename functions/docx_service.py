import base64
from io import BytesIO
from docxtpl import InlineImage
from docx.shared import Mm
from PIL import Image

def reparar_tags_rotos(documento_docx):
    def fusionar_parrafo(p):
        texto_completo = "".join(run.text for run in p.runs)
        if "{%" in texto_completo and "tr" in texto_completo:
            texto_corregido = texto_completo.replace("{% tr", "{%tr").replace("{%  tr", "{%tr")
            if len(p.runs) > 0:
                p.runs[0].text = texto_corregido
                for run in p.runs[1:]:
                    run.text = ""

    for p in documento_docx.paragraphs:
        fusionar_parrafo(p)
    for tabla in documento_docx.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    fusionar_parrafo(p)


def procesar_datos_rec(sub_contexto, doc):
    if isinstance(sub_contexto, dict):
        for k, v in list(sub_contexto.items()):
            if v is None:
                sub_contexto[k] = ""
            elif isinstance(v, str) and ("base64," in v or v.startswith("data:image")):
                try:
                    foto_str = v.split("base64,")[-1]
                    img_bytes_raw = base64.b64decode(foto_str)
                    pil_img = Image.open(BytesIO(img_bytes_raw)).convert("RGB")
                    buffer_corregido = BytesIO()
                    pil_img.save(buffer_corregido, format="JPEG", dpi=(96, 96))
                    buffer_corregido.seek(0)

                    if k == "FotoPerfil":
                        sub_contexto[k] = InlineImage(doc, buffer_corregido, width=Mm(30), height=Mm(40))
                    else:
                        sub_contexto[k] = InlineImage(doc, buffer_corregido, width=Mm(50))
                except Exception as e:
                    print(f"No se pudo procesar la imagen en '{k}': {e}")
                    sub_contexto[k] = ""
            else:
                procesar_datos_rec(v, doc)
    elif isinstance(sub_contexto, list):
        for i in range(len(sub_contexto)):
            if sub_contexto[i] is None:
                sub_contexto[i] = ""
            else:
                procesar_datos_rec(sub_contexto[i], doc)