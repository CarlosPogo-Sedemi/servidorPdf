"""Pinta la Matriz GERITRA (columnas A..AO) con openpyxl."""
from __future__ import annotations

import math
from datetime import date
from io import BytesIO
from itertools import groupby

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from functions import geritra_config as cfg
from functions.geritra_service import Fila, MatrizGeritra, Persona, construir

FUENTE = "Arial"
_fino = Side(style="thin", color="FF000000")
_medio = Side(style="medium", color="FF000000")
BORDE = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
L = Alignment(horizontal="left", vertical="center", wrap_text=True)
LT = Alignment(horizontal="left", vertical="top", wrap_text=True)

ULTIMA = "AO"
H1, H2 = 14, 15          # filas de encabezado de columnas
FILA_INI = 16             # primera fila de datos

ANCHOS = {
    "A": 9, "B": 16, "C": 18, "D": 24, "E": 14, "F": 16, "G": 18, "H": 30, "I": 26, "J": 34,
    "K": 30, "L": 10, "M": 10, "N": 12, "O": 12, "P": 11, "Q": 11, "R": 12, "S": 11, "T": 13, "U": 20,
    "V": 26, "W": 22, "X": 32, "Y": 36, "Z": 28, "AA": 26, "AB": 32, "AC": 30,
    "AD": 10, "AE": 12, "AF": 12, "AG": 11, "AH": 11, "AI": 12, "AJ": 11, "AK": 13,
    "AL": 26, "AM": 26, "AN": 13, "AO": 30,
}

ENCABEZADOS = [  # (rango, texto, sección)
    ("A", "NUMERO DE PROCESO", "ID"), ("B", "1. PROCESO", "ID"), ("C", "2. PUESTO DE TRABAJO", "ID"),
    ("D", "3. PERSONA CON DISCAPACIDAD", "ID"), ("E", "4. PORCENTAJE DE DISCAPACIDAD", "ID"),
    ("F", "5. AYUDAS TÉCNICAS EN EL TRABAJO", "ID"), ("G", "6. OBSERVACIÓN", "ID"),
    ("H", "7. ACTIVIDAD RUTINARIA", "ID"), ("I", "8. ACTIVIDAD NO RUTINARIA", "ID"),
    ("J", "9. MÁQUINAS, EQUIPOS, SUSTANCIAS Y ENERGÍAS EMPLEADAS", "ID"),
    ("K", "10. PELIGRO", "PELIGROS"), ("L", "11. CANTIDAD DE PERSONAS EXPUESTAS", "PELIGROS"),
    ("M", "12. ÍNDICE DE PERSONAS EXPUESTAS", "PELIGROS"), ("N", "13. ÍNDICE DE PROCEDIMIENTOS EXISTENTES", "PELIGROS"),
    ("O", "14. ÍNDICE DE CAPACITACIONES CAPACID. HUMANAS", "PELIGROS"), ("P", "15. ÍNDICE DE EXPOSICIÓN AL RIESGO", "PELIGROS"),
    ("Q", "16. PROBABILIDAD (P = Σ ÍNDICES)", "PELIGROS"), ("R", "17. CONSECUENCIAS (SEVERIDAD)", "PELIGROS"),
    ("S", "18. VALORACIÓN DEL RIESGO", "PELIGROS"), ("T", "19. NIVEL DE RIESGO", "PELIGROS"),
    ("U", "20. CLASIFICACIÓN DE LOS RIESGOS", "PELIGROS"),
    ("V", "21. ELIMINAR", "JERARQUIA"), ("W", "22. SUSTITUIR POR", "JERARQUIA"),
    ("X", "23. CONTROL DE INGENIERÍA Y DISEÑO", "JERARQUIA"),
    ("AD", "26. ÍNDICE DE PERSONAS EXPUESTAS", "RESIDUAL"), ("AE", "27. ÍNDICE DE PROCEDIMIENTOS EXISTENTES (2)", "RESIDUAL"),
    ("AF", "28. ÍNDICE DE CAPACITACIONES CAPACID. HUMANAS (2)", "RESIDUAL"), ("AG", "29. ÍNDICE DE EXPOSICIÓN AL RIESGO (2)", "RESIDUAL"),
    ("AH", "30. PROBABILIDAD (P = Σ ÍNDICES) (2)", "RESIDUAL"), ("AI", "31. CONSECUENCIAS (SEVERIDAD) (2)", "RESIDUAL"),
    ("AJ", "32. VALORACIÓN DEL RIESGO (2)", "RESIDUAL"), ("AK", "33. NIVEL DE RIESGO (2)", "RESIDUAL"),
    ("AL", "34. EFECTOS SOBRE LA SALUD DE LOS TRABAJADORES", "RESIDUAL"), ("AM", "35. VIGILANCIA DE LA SALUD", "RESIDUAL"),
    ("AN", "36. FECHA DE CONTROL E INSPECCIONES", "RESIDUAL"), ("AO", "37. ACTIVIDADES DE MEJORA CONTINUA", "RESIDUAL"),
]
SUBENCABEZADOS = [
    ("Y14:AA14", "24. CONTROL ADMINISTRATIVO"), ("AB14:AC14", "25. CONTROL EN EL TRABAJADOR"),
    ("Y15", "NORMAS, PLANES, PROCEDIMIENTOS, PERMISOS"), ("Z15", "ORGANIZACIÓN DE HORARIOS, DESCANSOS, ROTACIÓN DE PUESTOS"),
    ("AA15", "SEÑALIZACIÓN"), ("AB15", "FORMACIÓN, INFORMACIÓN Y CAPACITACIÓN"),
    ("AC15", "EQUIPOS DE PROTECCIÓN COLECTIVA E INDIVIDUAL"),
]
SECCIONES = [
    ("A13:J13", "IDENTIFICACIÓN DEL PUESTO DE TRABAJO", "ID"),
    ("K13:U13", "IDENTIFICACIÓN DE PELIGROS Y DIAGNÓSTICO INTEGRAL DE RIESGOS", "PELIGROS"),
    ("V13:AC13", "JERARQUÍA DE LOS CONTROLES EN LA SEGURIDAD Y SALUD EN EL TRABAJO", "JERARQUIA"),
    ("AD13:AO13", "GESTIÓN DEL RIESGO RESIDUAL", "RESIDUAL"),
]


def _fill(c: str) -> PatternFill:
    return PatternFill("solid", start_color=c, end_color=c)


def celda(ws: Worksheet, rango: str, valor=None, *, bold=False, size=10, color="FF000000",
          bg=None, al=C, formato=None, borde=True) -> None:
    min_c, min_r, max_c, max_r = range_boundaries(rango if ":" in rango else f"{rango}:{rango}")
    if (min_c, min_r) != (max_c, max_r):
        ws.merge_cells(start_row=min_r, start_column=min_c, end_row=max_r, end_column=max_c)
    p = ws.cell(row=min_r, column=min_c)
    if valor is not None:
        p.value = valor
    p.font = Font(name=FUENTE, size=size, bold=bold, color=color)
    p.alignment = al
    if formato:
        p.number_format = formato
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            cel = ws.cell(row=r, column=c)
            if borde:
                cel.border = BORDE
            if bg:
                cel.fill = _fill(bg)


def _lineas(txt: str, ancho: float, size: int = 9) -> int:
    if not txt:
        return 1
    chars = max(1, int(ancho * 1.15 * 10 / size))
    return sum(max(1, math.ceil(len(p) / chars)) for p in str(txt).split("\n"))


def _vinetas(items: list[str]) -> str:
    return "\n".join(f"• {x}" for x in items)


# ----------------------------------------------------------------------------
# Discapacidad (opción A): solo personas con discapacidad que hacen la actividad
# ----------------------------------------------------------------------------
def _discapacidad(f: Fila) -> tuple[str, str, str, str]:
    if not f.disc_personas:
        return ("NO", "N/A", "N/A", "N/A")
    ps = f.disc_personas
    d = f"SÍ ({len(ps)} de {f.total_personas_actividad})\n" + "\n".join(f"• {p.nombre}: {p.tipo}" for p in ps)
    pct = lambda v: v if (not v or "%" in v or not v.replace(".", "").isdigit()) else f"{v} %"
    e = "\n".join(f"• {p.nombre}: {pct(p.porcentaje) or '—'}" for p in ps)
    fa = "\n".join(f"• {p.nombre}: {p.ayuda or '—'}" for p in ps)
    g = "\n".join(f"• {p.nombre}: {p.observacion or '—'}" for p in ps)
    return d, e, fa, g


# ----------------------------------------------------------------------------
# Secciones
# ----------------------------------------------------------------------------
def _encabezado(ws: Worksheet, m: MatrizGeritra) -> None:
    ws.row_dimensions[1].height = 60
    celda(ws, f"A1:{ULTIMA}1", cfg.TITULO, bold=True, size=20, color=cfg.COLOR_AZUL)
    celda(ws, f"A2:{ULTIMA}2", cfg.SUBTITULO, bold=True, size=11, color="FFFFFFFF", bg=cfg.COLOR_AZUL)
    fmt = "DD/MM/YYYY"
    izq = [
        ("RAZÓN SOCIAL", m.razon_social), ("RUC", m.ruc), ("DIRECCIÓN", m.direccion),
        ("TELÉFONO", m.telefono), ("RESPONSABLE DE LA EVALUACIÓN", m.responsable_evaluacion),
        ("RESPONSABLE DE SSO", m.responsable_sso), ("FECHA DE LA EVALUACIÓN", m.fecha_evaluacion),
    ]
    for i, (lbl, val) in enumerate(izq):
        r = 3 + i
        ws.row_dimensions[r].height = 22
        celda(ws, f"A{r}:C{r}", lbl, bold=True, bg=cfg.COLOR_GRIS, al=L)
        celda(ws, f"D{r}:J{r}", val, al=L, formato=fmt if isinstance(val, date) else None)
    der = [
        ("CÓDIGO AST", m.codigo_ast), ("PROCESO", m.proceso), ("PUESTO DE TRABAJO", m.puesto),
        ("TOTAL DE TRABAJADORES DEL PUESTO", m.numero_trabajadores),
        ("PERSONAS ENTREVISTADAS (MUESTRA)", m.personas_entrevistadas),
        ("PERSONAS CON DISCAPACIDAD", len(m.personas_discapacidad)),
    ]
    for i, (lbl, val) in enumerate(der):
        r = 3 + i
        celda(ws, f"L{r}:P{r}", lbl, bold=True, bg=cfg.COLOR_GRIS, al=L)
        celda(ws, f"Q{r}:U{r}", val, al=L)
    # Leyenda: matriz de niveles de riesgo (Probabilidad x Consecuencias) + acción por nivel
    blanco = "FFFFFFFF"
    celda(ws, "W3:AF3", "NIVELES DE RIESGO", bold=True, size=12, color=blanco, bg=cfg.COLOR_AZUL)
    celda(ws, "W4:X5", "VALORACIÓN =\nPROBABILIDAD × CONSECUENCIAS", bold=True, size=8, bg=cfg.COLOR_GRIS)
    celda(ws, "Y4:AA4", "CONSECUENCIAS", bold=True, size=10, bg=cfg.COLOR_GRIS)
    for col, txt in zip(("Y", "Z", "AA"), ("Ligeramente dañino\nLD", "Dañino\nD", "Extremadamente dañino\nED")):
        celda(ws, f"{col}5", txt, bold=True, size=9, bg=cfg.COLOR_GRIS)
    celda(ws, "W6:W8", "PROBABILIDAD", bold=True, size=10, bg=cfg.COLOR_GRIS,
          al=Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True))
    for i, (prob, fila) in enumerate(zip(("Baja\nB", "Media\nM", "Alta\nA"), cfg.MATRIZ_NIVELES)):
        r = 6 + i
        celda(ws, f"X{r}", prob, bold=True, size=9, bg=cfg.COLOR_GRIS)
        for col, nivel in zip(("Y", "Z", "AA"), fila):
            celda(ws, f"{col}{r}", f"Riesgo {nivel.lower()}\n{cfg.CODIGO_NIVEL[nivel]}", bold=True, size=9,
                  bg=cfg.COLOR_NIVEL[nivel])
    celda(ws, "W9:AA9", cfg.NOTA_PROBABILIDAD, size=7, al=L)
    celda(ws, "AB4:AF4", "ACCIÓN Y TEMPORIZACIÓN", bold=True, size=10, bg=cfg.COLOR_GRIS)
    for i, nivel in enumerate(cfg.COLOR_NIVEL):
        r = 5 + i
        celda(ws, f"AB{r}", f"{nivel.upper()} ({cfg.CODIGO_NIVEL[nivel]})", bold=True, size=9, bg=cfg.COLOR_NIVEL[nivel])
        celda(ws, f"AC{r}:AF{r}", cfg.ACCION_NIVEL[nivel], size=8, al=L)
    for r in range(4, 10):
        ws.row_dimensions[r].height = 34

    ws.row_dimensions[13].height = 28
    for rng, txt, sec in SECCIONES:
        celda(ws, rng, txt, bold=True, size=11, bg=cfg.COLOR_SECCION[sec])
    ws.row_dimensions[H1].height = 30
    ws.row_dimensions[H2].height = 62
    for col, txt, sec in ENCABEZADOS:
        celda(ws, f"{col}{H1}:{col}{H2}", txt, bold=True, size=8, bg=cfg.COLOR_SECCION[sec])
    for rng, txt in SUBENCABEZADOS:
        celda(ws, rng, txt, bold=True, size=8, bg=cfg.COLOR_SECCION["JERARQUIA"])


def _valor_nivel(ws, col: str, r: int, nivel: str) -> None:
    celda(ws, f"{col}{r}", nivel or None, bold=True, size=9, bg=cfg.COLOR_NIVEL.get(nivel))


def _fila(ws: Worksheet, r: int, f: Fila) -> None:
    ini, res = f.inicial, f.residual
    peligro = f.factor + (f"\n(Peligro: {f.peligro})" if f.peligro and f.peligro.lower() != f.factor.lower() else "")
    celda(ws, f"K{r}", peligro, size=9, al=L)
    for col, val in zip("LMNOPQRS", [ini.cantidad, ini.i_personas, ini.i_procedimientos, ini.i_capacitaciones,
                                       ini.i_exposicion, ini.probabilidad, ini.i_consecuencias, ini.valoracion]):
        celda(ws, f"{col}{r}", val, size=10, bold=col in "QS")
    _valor_nivel(ws, "T", r, ini.nivel)
    celda(ws, f"U{r}", f.clasificacion, size=8)

    for clave, col in cfg.JERARQUIA_COLUMNAS:
        celda(ws, f"{col}{r}", _vinetas(f.jerarquia.get(clave, [])) or None, size=8, al=LT)

    if res:
        for col, val in zip(["AD", "AE", "AF", "AG", "AH", "AI", "AJ"],
                            [res.i_personas, res.i_procedimientos, res.i_capacitaciones, res.i_exposicion,
                             res.probabilidad, res.i_consecuencias, res.valoracion]):
            celda(ws, f"{col}{r}", val, size=10, bold=col in ("AH", "AJ"))
        _valor_nivel(ws, "AK", r, res.nivel)
    else:
        for col in ["AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK"]:
            celda(ws, f"{col}{r}", None)
    celda(ws, f"AL{r}", f.efectos or None, size=8, al=L)
    celda(ws, f"AM{r}", f.vigilancia or None, size=8, al=L)
    celda(ws, f"AN{r}", f.fecha_control, size=9, formato="DD/MM/YYYY")
    celda(ws, f"AO{r}", f.mejora or None, size=8, al=L)

    alto = max([_lineas(peligro, ANCHOS["K"])] +
               [_lineas(_vinetas(f.jerarquia.get(k, [])), ANCHOS[c], 8) for k, c in cfg.JERARQUIA_COLUMNAS] +
               [_lineas(t, ANCHOS[c], 8) for t, c in ((f.efectos, "AL"), (f.vigilancia, "AM"), (f.mejora, "AO"))] +
               [2])
    ws.row_dimensions[r].height = min(409, alto * 11 + 8)


def _asegurar_alto(ws: Worksheet, ini: int, fin: int, minimo: float) -> None:
    actual = sum(ws.row_dimensions[i].height or 15 for i in range(ini, fin + 1))
    if actual < minimo:
        extra = (minimo - actual) / (fin - ini + 1)
        for i in range(ini, fin + 1):
            ws.row_dimensions[i].height = (ws.row_dimensions[i].height or 15) + extra


def _cuerpo(ws: Worksheet, m: MatrizGeritra) -> int:
    r = FILA_INI
    for _, grupo in groupby(m.filas, key=lambda f: f.orden[0]):
        filas = list(grupo)
        ini = r
        for f in filas:
            _fila(ws, r, f)
            r += 1
        fin = r - 1
        d, e, fa, g = _discapacidad(filas[0])
        celda(ws, f"H{ini}:H{fin}", filas[0].actividad, size=9, al=L, bold=True)
        celda(ws, f"D{ini}:D{fin}", d, size=8, al=L)
        celda(ws, f"E{ini}:E{fin}", e, size=8, al=L)
        celda(ws, f"F{ini}:F{fin}", fa, size=8, al=L)
        celda(ws, f"G{ini}:G{fin}", g, size=8, al=L)
        _asegurar_alto(ws, ini, fin, max(_lineas(d, ANCHOS["D"], 8), _lineas(e, ANCHOS["E"], 8),
                                         _lineas(fa, ANCHOS["F"], 8), _lineas(g, ANCHOS["G"], 8),
                                         _lineas(filas[0].actividad, ANCHOS["H"])) * 11 + 10)
    fin = r - 1
    maq = "\n\n".join(f"{grupo}:\n" + _vinetas(items) for grupo, items in m.maquinas.items()) or "N/A"
    celda(ws, f"A{FILA_INI}:A{fin}", 1, bold=True, size=12)   # consecutivo del proceso
    celda(ws, f"B{FILA_INI}:B{fin}", m.proceso, bold=True, size=10)
    celda(ws, f"C{FILA_INI}:C{fin}", m.puesto, bold=True, size=10)
    celda(ws, f"I{FILA_INI}:I{fin}", _vinetas(m.no_rutinarias) or "N/A", size=8, al=LT)
    celda(ws, f"J{FILA_INI}:J{fin}", maq, size=8, al=LT)
    _asegurar_alto(ws, FILA_INI, fin, max(_lineas(maq, ANCHOS["J"], 8), _lineas(_vinetas(m.no_rutinarias), ANCHOS["I"], 8)) * 11)
    return r


def _formula_nivel(ref: str) -> str:
    return (f'=IF({ref}="","",IF({ref}<5,"Trivial",IF({ref}<9,"Tolerable",'
            f'IF({ref}<17,"Moderado",IF({ref}<25,"Importante","Intolerable")))))')


def _formula_accion(ref: str) -> str:
    f = '""'
    for nivel, txt in reversed(list(cfg.ACCION_NIVEL.items())):
        f = f'IF({ref}="{nivel}","{txt}",{f})'
    return "=" + f


def _promedios(ws: Worksheet, r: int, fin_datos: int) -> int:
    """Fila resumen del puesto: promedio de valoración inicial y residual (como el VBA)."""
    ws.row_dimensions[r].height = 48
    celda(ws, f"A{r}:R{r}", "PROMEDIO DEL PUESTO DE TRABAJO (valoración inicial y residual)", bold=True,
          size=10, bg=cfg.COLOR_GRIS, al=Alignment(horizontal="right", vertical="center", wrap_text=True))
    rng_s = f"S{FILA_INI}:S{fin_datos}"
    rng_aj = f"AJ{FILA_INI}:AJ{fin_datos}"
    celda(ws, f"S{r}", f'=IFERROR(ROUND(AVERAGE({rng_s}),2),"")', bold=True, formato="0.00")
    celda(ws, f"T{r}", _formula_nivel(f"S{r}"), bold=True)
    celda(ws, f"U{r}:AC{r}", _formula_accion(f"T{r}"), size=8, al=L)
    celda(ws, f"AD{r}:AI{r}", "PROMEDIO RESIDUAL", bold=True, size=9, bg=cfg.COLOR_GRIS)
    celda(ws, f"AJ{r}", f'=IFERROR(ROUND(AVERAGE({rng_aj}),2),"")', bold=True, formato="0.00")
    celda(ws, f"AK{r}", _formula_nivel(f"AJ{r}"), bold=True)
    celda(ws, f"AL{r}:AO{r}", _formula_accion(f"AK{r}"), size=8, al=L)
    return r


def _hoja_discapacidad(wb: Workbook, m: MatrizGeritra) -> None:
    ws = wb.create_sheet("PERSONAS CON DISCAPACIDAD")
    ws.sheet_view.showGridLines = False
    anchos = {"A": 5, "B": 32, "C": 26, "D": 22, "E": 12, "F": 26, "G": 30, "H": 20, "I": 55, "J": 45}
    for k, v in anchos.items():
        ws.column_dimensions[k].width = v
    celda(ws, "A1:J1", f"PERSONAS CON DISCAPACIDAD — {m.puesto} ({m.codigo_ast})", bold=True, size=13,
          color="FFFFFFFF", bg=cfg.COLOR_AZUL)
    ws.row_dimensions[1].height = 28
    celda(ws, "A2:J2", "Información sensible de uso interno de SSO. Detalle del resumen de las columnas 3 a 6 de la matriz.",
          size=8, al=L, borde=False)
    cab = ["N°", "NOMBRE", "CARGO", "TIPO DE DISCAPACIDAD", "PORCENTAJE", "AYUDAS TÉCNICAS",
           "OBSERVACIÓN", "CATEGORÍA DE VULNERABILIDAD", "ACTIVIDADES RUTINARIAS",
           "ACTIVIDADES NO RUTINARIAS"]
    for col, txt in zip("ABCDEFGHIJ", cab):
        celda(ws, f"{col}4", txt, bold=True, size=9, bg=cfg.COLOR_GRIS)
    ws.row_dimensions[4].height = 30
    if not m.personas_discapacidad:
        celda(ws, "A5:J5", "No se registran personas con discapacidad en las entrevistas de este puesto.", size=9)
        return
    for i, p in enumerate(m.personas_discapacidad, 1):
        r = 4 + i
        acts = _vinetas(p.actividades_txt) or "N/A"
        no_rut = _vinetas(list(dict.fromkeys(p.eventos))) or "N/A"
        for col, val in zip("ABCDEFGHIJ", [i, p.nombre, p.cargo, p.tipo, p.porcentaje, p.ayuda,
                                            p.observacion, p.categoria, acts, no_rut]):
            celda(ws, f"{col}{r}", val, size=9, al=L if col in "BCFGIJ" else C)
        ws.row_dimensions[r].height = max(30, max(_lineas(acts, anchos["I"]), _lineas(no_rut, anchos["J"])) * 12 + 6)


def _logos(ws: Worksheet) -> None:
    for archivo, ancla, alto in (("logo_sedemi.png", "A1", 40), ("logo_mrl.png", "AL1", 60)):
        ruta = cfg.ASSETS_DIR / archivo
        if ruta.exists():
            img = XLImage(str(ruta))
            esc = alto / img.height
            img.height, img.width = alto, int(img.width * esc)
            ws.add_image(img, ancla)


def _impresion(ws: Worksheet, ultima: int) -> None:
    ws.print_title_rows = f"{H1}:{H2}"
    ws.print_area = f"A1:{ULTIMA}{ultima}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.zoomScale = 70
    ws.sheet_view.showGridLines = False


# ----------------------------------------------------------------------------
# API pública
# ----------------------------------------------------------------------------
def generar_excel_geritra(data: dict) -> tuple[bytes, str]:
    m = construir(data)
    wb = Workbook()
    ws = wb.active
    ws.title = "MATRIZ GERITRA"
    for col, ancho in ANCHOS.items():
        ws.column_dimensions[col].width = ancho

    _encabezado(ws, m)
    if m.filas:
        siguiente = _cuerpo(ws, m)
        ultima = _promedios(ws, siguiente, siguiente - 1)
    else:
        celda(ws, f"A{FILA_INI}:{ULTIMA}{FILA_INI}", "El AST no tiene registros en la matriz de riesgos.", bold=True)
        ultima = FILA_INI
    _logos(ws)
    _impresion(ws, ultima)
    _hoja_discapacidad(wb, m)
    wb.calculation.fullCalcOnLoad = True

    buffer = BytesIO()
    wb.save(buffer)
    puesto = "".join(ch if ch.isalnum() else "_" for ch in m.puesto)[:60].strip("_")
    return buffer.getvalue(), f"GERITRA_{m.codigo_ast}_{puesto}.xlsx"
